"""Grades, results computation and CSV/Excel import with validation."""
from __future__ import annotations

import csv
import io

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import audit
from .db import utcnow
from .models import (
    AssessmentComponent,
    Course,
    CourseResult,
    Enrollment,
    GradeChange,
    GradeEntry,
    User,
)
from .tenancy import letter_for

RESULT_FLOW = ["draft", "submitted", "approved", "published"]


def result_for(db: Session, course_id: int) -> CourseResult:
    res = db.scalar(select(CourseResult).where(CourseResult.course_id == course_id))
    if not res:
        res = CourseResult(course_id=course_id, status="draft")
        db.add(res)
        db.flush()
    return res


def components(db: Session, course_id: int) -> list[AssessmentComponent]:
    return list(db.scalars(select(AssessmentComponent).where(AssessmentComponent.course_id == course_id)
                           .order_by(AssessmentComponent.position, AssessmentComponent.id)))


def student_summary(db: Session, course: Course, student_id: int, grading: dict,
                    comps: list[AssessmentComponent] | None = None) -> dict:
    comps = comps if comps is not None else components(db, course.id)
    entries = {e.component_id: e for e in db.scalars(select(GradeEntry).where(
        GradeEntry.course_id == course.id, GradeEntry.student_id == student_id))}
    scores = {}
    total = 0.0
    max_total = 0.0
    complete = True
    for c in comps:
        e = entries.get(c.id)
        score = e.score if e else None
        scores[str(c.id)] = score
        max_total += c.max_score
        if score is None:
            complete = False
        else:
            total += score
    percent = round(100 * total / max_total, 1) if max_total else 0.0
    letter, label = letter_for(percent, grading) if complete and comps else (None, None)
    return {"scores": scores, "total": round(total, 2), "max_total": max_total, "percent": percent,
            "complete": complete and bool(comps), "letter": letter, "label": label,
            "passed": (percent >= grading["pass_mark"]) if complete and comps else None}


def set_score(db: Session, actor: User, course: Course, student_id: int, component: AssessmentComponent,
              score: float | None, reason: str, *, source: str = "manual", ip: str | None = None) -> bool:
    """Set one score; returns True if it changed. Every change is logged with its reason."""
    if score is not None and not (0 <= score <= component.max_score):
        raise HTTPException(status_code=400, detail=f"الدرجة يجب أن تكون بين 0 و {component.max_score}")
    entry = db.scalar(select(GradeEntry).where(GradeEntry.student_id == student_id,
                                               GradeEntry.component_id == component.id))
    old = entry.score if entry else None
    if old == score:
        return False
    if old is not None and (not reason or len(reason.strip()) < 3):
        raise HTTPException(status_code=400, detail="تعديل درجة مسجلة يتطلب ذكر السبب")
    if not entry:
        entry = GradeEntry(course_id=course.id, student_id=student_id, component_id=component.id)
        db.add(entry)
    entry.score, entry.updated_by, entry.updated_at = score, actor.id, utcnow()
    db.add(GradeChange(course_id=course.id, student_id=student_id, component_id=component.id, old_score=old,
                       new_score=score, reason=(reason or "إدخال أولي").strip(), source=source, actor_id=actor.id))
    if old is not None:
        audit.record(db, "grade.changed", actor=actor, target_type="grade", target_id=f"{course.id}:{student_id}",
                     details={"component": component.name, "old": old, "new": score, "reason": reason,
                              "source": source}, ip=ip)
    return True


# --------------------------------------------------------------------------- import


def _rows_from_upload(filename: str, data: bytes) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=400, detail="يجب حفظ ملف CSV بترميز UTF-8") from exc
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        from .files import FileRejected, inspect

        try:
            inspect(data, "xlsx")
        except FileRejected as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise HTTPException(status_code=400, detail="الصيغ المدعومة: CSV أو XLSX")


def _num(value: str) -> float | None:
    value = (value or "").strip().replace(",", ".")
    if value == "":
        return None
    # accept Arabic-Indic digits
    value = value.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    return float(value)


def parse_import(db: Session, course: Course, filename: str, data: bytes) -> tuple[list, list]:
    """Validate a grade sheet. Expected header: student_number, <component name>, ..."""
    rows = [r for r in _rows_from_upload(filename, data) if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="الملف لا يحتوي على بيانات")
    if len(rows) > 5000:
        raise HTTPException(status_code=400, detail="عدد الصفوف كبير جداً")
    header = [h.strip() for h in rows[0]]
    comps = {c.name: c for c in components(db, course.id)}
    errors: list[dict] = []
    if not header or header[0].lower() not in ("student_number", "رقم القيد", "student", "code"):
        errors.append({"row": 1, "error": "العمود الأول يجب أن يكون رقم القيد (student_number)"})
        return [], errors
    col_map = {}
    for idx, name in enumerate(header[1:], start=1):
        if name == "name" or name == "الاسم":
            continue
        if name not in comps:
            errors.append({"row": 1, "error": f"عنصر تقييم غير معروف: {name}"})
        else:
            col_map[idx] = comps[name]
    enrolled = {
        u.code: u for u in db.scalars(
            select(User).join(Enrollment, Enrollment.student_id == User.id)
            .where(Enrollment.course_id == course.id, Enrollment.status == "active"))
    }
    parsed = []
    seen = set()
    for n, row in enumerate(rows[1:], start=2):
        code = (row[0] if row else "").strip().translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
        student = enrolled.get(code)
        if not student:
            errors.append({"row": n, "error": f"رقم القيد {code or '(فارغ)'} غير مسجل في هذا المقرر"})
            continue
        if code in seen:
            errors.append({"row": n, "error": f"رقم القيد {code} مكرر"})
            continue
        seen.add(code)
        scores = {}
        for idx, comp in col_map.items():
            raw = row[idx] if idx < len(row) else ""
            try:
                value = _num(raw)
            except ValueError:
                errors.append({"row": n, "error": f"قيمة غير رقمية في {comp.name}: {raw}"})
                continue
            if value is not None and not (0 <= value <= comp.max_score):
                errors.append({"row": n, "error": f"{comp.name}: {value} خارج المدى 0-{comp.max_score}"})
                continue
            scores[str(comp.id)] = value
        parsed.append({"student_id": student.id, "code": code, "name": student.full_name, "scores": scores})
    return parsed, errors
